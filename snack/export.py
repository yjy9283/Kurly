"""엑셀 생성: 구매목록 / 카테고리요약 / 후보전체. 금액은 수식으로 계산."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from . import config as C

HEAD = PatternFill("solid", fgColor="5F0080")  # 컬리 퍼플
THIN = Side(style="thin", color="DDDDDD")


def _header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c)
        x.font = Font(bold=True, color="FFFFFF"); x.fill = HEAD
        x.alignment = Alignment(horizontal="center", vertical="center")


def export(picks, cands, warnings, path):
    wb = Workbook()
    ws = wb.active; ws.title = "구매목록"
    ws["A1"] = "회사 다과 구매 목록 (마켓컬리)"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "예산"; ws["B2"] = C.BUDGET; ws["B2"].number_format = "#,##0"
    ws["C2"] = "합계"; ws["D2"] = None
    ws["E2"] = "잔액"; ws["F2"] = None
    cols = ["구분", "카테고리", "상품명", "단가(원)", "수량", "금액(원)", "할인율(%)", "후기수", "링크", "메모"]
    _header(ws, 4, cols)
    picks = sorted(picks, key=lambda p: (not p["must"], p["category"]))
    for r, p in enumerate(picks, 5):
        ws.append([("필수" if p["must"] else "추천"), p["category"], p["name"], p["price"], p["qty"],
                   f"=D{r}*E{r}", p["discount_rate"], p["reviews"], p["url"], ""])
        ws.cell(row=r, column=9).hyperlink = p["url"]; ws.cell(row=r, column=9).font = Font(color="0563C1", underline="single")
    last = 4 + len(picks)
    ws["D2"] = f"=SUM(F5:F{last})"; ws["F2"] = f"=B2-D2"
    for c in ("D2", "F2"): ws[c].number_format = "#,##0"; ws[c].font = Font(bold=True)
    for r in range(5, last + 1):
        for c in (4, 6): ws.cell(row=r, column=c).number_format = "#,##0"
        if ws.cell(row=r, column=1).value == "필수":
            for c in range(1, 11): ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="FFF2CC")
    for i, w in enumerate([7, 14, 52, 11, 7, 12, 10, 9, 36, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    if warnings:
        ws.cell(row=last + 2, column=1, value="⚠ " + " / ".join(warnings)).font = Font(color="C00000", bold=True)

    s = wb.create_sheet("카테고리요약")
    _header(s, 1, ["카테고리", "품목수", "금액(원)", "비중"])
    cats = ["필수"] + list(C.CATEGORIES)
    for r, cat in enumerate(cats, 2):
        s.append([cat, f'=COUNTIF(구매목록!B5:B{last},A{r})', f'=SUMIF(구매목록!B5:B{last},A{r},구매목록!F5:F{last})',
                  f"=C{r}/SUM($C$2:$C${len(cats)+1})"])
        s.cell(row=r, column=3).number_format = "#,##0"; s.cell(row=r, column=4).number_format = "0.0%"
    for i, w in enumerate([18, 8, 14, 8], 1): s.column_dimensions[get_column_letter(i)].width = w

    a = wb.create_sheet("후보전체")
    _header(a, 1, ["상품번호", "상품명", "가격", "정가", "할인율", "후기수", "품절", "검색어", "링크"])
    for p in sorted(cands, key=lambda p: p["keyword"]):
        a.append([p["no"], p["name"], p["price"], p["list_price"], p["discount_rate"], p["reviews"], p["sold_out"], p["keyword"], p["url"]])
    a.auto_filter.ref = a.dimensions
    for i, w in enumerate([10, 60, 10, 10, 8, 8, 6, 16, 36], 1): a.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)
